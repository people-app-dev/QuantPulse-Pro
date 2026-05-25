"""
Template inventory & output validation helper for research-to-slides skill.
Usage:
  python template_helper.py inventory <template.pptx>   — Print slide inventory
  python template_helper.py validate <output.pptx>       — Check for placeholder text
"""
import sys
from pptx import Presentation


def inventory(path):
    """Print a structured inventory of all slides in the template."""
    prs = Presentation(path)
    print(f"Template: {path}")
    print(f"Slides: {len(prs.slides)} | Size: {prs.slide_width}x{prs.slide_height}")
    print(f"Layouts: {len(prs.slide_layouts)}")
    print("=" * 60)

    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} (Layout: {slide.slide_layout.name}) ---")
        texts = []
        tables = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append((shape.name, t[:80]))
            if shape.has_table:
                tbl = shape.table
                rows = len(tbl.rows)
                cols = len(tbl.columns)
                # Get header row
                header = [tbl.cell(0, c).text.strip()[:20] for c in range(min(cols, 6))]
                tables.append((shape.name, rows, cols, header))

        if texts:
            for name, t in texts:
                print(f"  TXT [{name}]: {t}")
        if tables:
            for name, rows, cols, header in tables:
                print(f"  TBL [{name}]: {rows}r x {cols}c | headers: {' | '.join(header)}")


def validate(path):
    """Check for leftover placeholder text in output."""
    prs = Presentation(path)
    placeholders = ["XXX", "XXXX", "XXXXXX", "Lorem", "ipsum", "placeholder"]
    issues = []

    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        for ph in placeholders:
                            if ph.upper() in text.upper():
                                issues.append((i + 1, shape.name, text))

            if shape.has_table:
                for ri, row in enumerate(shape.table.rows):
                    for ci, cell in enumerate(row.cells):
                        text = cell.text.strip()
                        if text:
                            for ph in placeholders:
                                if ph.upper() in text.upper():
                                    issues.append((i + 1, f"{shape.name}[{ri},{ci}]", text))

    if issues:
        print(f"FOUND {len(issues)} PLACEHOLDER ISSUES:")
        for slide_num, shape_name, text in issues:
            print(f"  Slide {slide_num} [{shape_name}]: {text}")
    else:
        print("No placeholder text found — looks clean!")

    return len(issues)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python template_helper.py <inventory|validate> <file.pptx>")
        sys.exit(1)

    cmd = sys.argv[1]
    path = sys.argv[2]

    if cmd == "inventory":
        inventory(path)
    elif cmd == "validate":
        issues = validate(path)
        sys.exit(0 if issues == 0 else 1)
    else:
        print(f"Unknown command: {cmd}")
        sys.exit(1)
